#!/usr/bin/env python3
"""Generate the workshop PowerPoint decks from this script.

    sudo apt-get install -y python3-pptx    # Debian/Ubuntu
    python3 docs/workshop/build_decks.py    # writes the .pptx files next to this script

Two decks:
  1. workshop-demo.pptx      - the talk track for a team tech workshop
  2. technical-deep-dive.pptx - architecture, decisions, adoption (reference)

Content is generic (this public reference implementation). Add workplace
specifics verbally on the day.
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = Path(__file__).resolve().parent

INK = RGBColor(0x1F, 0x2A, 0x37)
ACCENT = RGBColor(0x43, 0x38, 0xCA)
MUTED = RGBColor(0x64, 0x70, 0x84)
BG = RGBColor(0xF4, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def title_slide(prs, title, subtitle, footer):
    s = _blank(prs)
    _bg(s, INK)
    bar = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6))
    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = WHITE
    sub = s.shapes.add_textbox(Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.0))
    p = sub.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = subtitle
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
    ft = s.shapes.add_textbox(Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5))
    p = ft.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = footer
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    return s


def _heading(slide, text):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = INK
    line = slide.shapes.add_shape(1, Inches(0.72), Inches(1.32), Inches(1.6), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def bullets_slide(prs, heading, bullets, note=None):
    s = _blank(prs)
    _bg(s)
    _heading(s, heading)
    body = s.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.0))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for level, text in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        r = p.add_run()
        r.text = text
        r.font.size = Pt(22 - 3 * level)
        r.font.color.rgb = INK if level == 0 else MUTED
        p.space_after = Pt(10)
    if note:
        nb = s.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.6))
        p = nb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = note
        r.font.size = Pt(13)
        r.font.italic = True
        r.font.color.rgb = MUTED
    return s


def demo_slide(prs, heading, steps, expect):
    s = _blank(prs)
    _bg(s)
    _heading(s, "DEMO  ·  " + heading)
    body = s.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.2))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for i, step in enumerate(steps, 1):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = f"{i}.  {step}"
        r.font.size = Pt(20)
        r.font.color.rgb = INK
        p.space_after = Pt(12)
    card = s.shapes.add_shape(1, Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.15))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xFF)
    card.line.color.rgb = ACCENT
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Expect:  " + expect
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)
    return s


def section_slide(prs, text):
    s = _blank(prs)
    _bg(s, ACCENT)
    box = s.shapes.add_textbox(Inches(0.9), Inches(3.1), Inches(11.5), Inches(1.4))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return s


def diagram_slide(prs, heading, rows, caption=None):
    """rows: list of strings, each a stage in the pipeline."""
    s = _blank(prs)
    _bg(s)
    _heading(s, heading)
    y = Inches(1.9)
    for i, row in enumerate(rows):
        card = s.shapes.add_shape(1, Inches(1.2), y, Inches(10.9), Inches(0.8))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(0xEE, 0xF2, 0xFF)
        card.line.color.rgb = RGBColor(0xE3, 0xE8, 0xEF)
        tf = card.text_frame
        tf.margin_left = Inches(0.25)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = row
        r.font.size = Pt(16)
        r.font.color.rgb = INK
        y += Inches(0.95)
        if i < len(rows) - 1:
            arr = s.shapes.add_textbox(Inches(6.4), y - Inches(0.28), Inches(0.6), Inches(0.3))
            pr = arr.text_frame.paragraphs[0]
            pr.alignment = PP_ALIGN.CENTER
            rr = pr.add_run()
            rr.text = "↓"
            rr.font.size = Pt(14)
            rr.font.color.rgb = MUTED
    if caption:
        cb = s.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.5))
        p = cb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = caption
        r.font.size = Pt(13)
        r.font.italic = True
        r.font.color.rgb = MUTED
    return s


# ---------------------------------------------------------------- deck 1
def build_demo_deck():
    prs = new_deck()
    title_slide(
        prs,
        "Load-Test Metrics for a JBoss / JSF App",
        "Getting user-action and app-server timings into Prometheus & Grafana — during a load test",
        "Team tech workshop  ·  reference implementation: github.com/mmcnicol/jboss-prometheus-grafana",
    )

    bullets_slide(prs, "The problem", [
        (0, "We run load tests before a release. We want to see:"),
        (1, "how long key user actions take (login → next page, open form, save)"),
        (1, "whether a new release is slower than the last one"),
        (1, "what the JVM and app server are doing under load"),
        (0, "Constraints"),
        (1, "load-test only — the app must not emit metrics all the time"),
        (1, "keep it simple; minimal change to the application"),
        (1, "Test / UAT / Prod are Windows; an incumbent APM agent is already there"),
    ])

    bullets_slide(prs, "What this is", [
        (0, "A small, portable metrics module + a run pipeline + Grafana dashboards"),
        (1, "one Maven module you drop into a WAR — no framework, no container coupling"),
        (1, "a JBoss system property turns it on/off:  -Dportal.metrics.enabled=true"),
        (1, "an OpenTelemetry Collector runs only for the duration of a load test"),
        (1, "each run tagged with run_id + release, so runs can be compared"),
        (0, "Demonstrated end-to-end on WildFly 26.1 (the free stand-in for EAP 7.4)"),
    ], note="Same JDK (17), same Jakarta EE 8 / javax namespace as EAP 7.4.")

    diagram_slide(prs, "How a run flows", [
        "Scenario driver (Selenium / k6 browser / k6 HTTP)  —  performs the steps",
        "WildFly / EAP  —  portal-web + services; server-side timer per action; /metrics served only when enabled",
        "OpenTelemetry Collector  —  started & stopped by the CI job; stamps run_id / release / test_type",
        "Prometheus  —  stores the run; recording rules roll up per-run p95",
        "Grafana  —  User Actions · Baseline vs Candidate · Release Trend · App Server & JVM",
    ], caption="The driver just drives. The authoritative timing is measured server-side.")

    section_slide(prs, "Demo")

    demo_slide(prs, "1 · The toggle", [
        "Show WildFly running with portal.metrics.enabled=false",
        "curl /portal-web/metrics  →  404 (no endpoint, no overhead)",
        "Flip the system property to true, restart, curl again",
    ], "404 when off; Prometheus-format metrics when on. Nothing else changes.")

    demo_slide(prs, "2 · A labelled run", [
        "./load/run-loadtest.sh --driver selenium --vus 3 --duration 120 --release 1.4.0",
        "It enables metrics, starts the Collector, runs the scenario, stops the Collector, restores state",
        "Open Grafana → 'User Actions — Load Test', pick the run",
    ], "Per-action throughput, error rate, p50/p90/p95/p99 over the run.")

    demo_slide(prs, "3 · Baseline vs candidate", [
        "Run the scenario again with a different --release",
        "Open 'Baseline vs Candidate', pick the two runs",
        "Read the bar chart + delta table (Δ p95 per action, colour-graded)",
    ], "A regression between releases is obvious. No manual spreadsheet.")

    demo_slide(prs, "4 · App server & JVM", [
        "Same run — the Collector also scraped WildFly's MicroProfile Metrics endpoint",
        "Open 'App Server & JVM — Load Test'",
        "Heap, GC pause, threads, CPU, Undertow requests/s — for that run only",
    ], "No agent installed. No application code. The subsystem is already in EAP 7.4.")

    demo_slide(prs, "5 · Service endpoints", [
        "./load/run-loadtest.sh --driver k6-http --release 1.4.0",
        "Open 'Service Endpoints — Load Test'",
        "Per-endpoint p95, error ratio, client (k6) vs server latency",
    ], "The 'middle' layer — real URLs, so k6's own metrics are appropriate here too.")

    section_slide(prs, "Why it stays simple")

    bullets_slide(prs, "Design choices that keep it small", [
        (0, "One facade, swappable backend"),
        (1, "app code calls Metrics.get().action(\"login\").record(elapsed, outcome)"),
        (1, "Micrometer by default; Prometheus Java client proven interchangeable"),
        (0, "Action name resolved server-side from JSF (no driver cooperation needed)"),
        (1, "works with Selenium, k6 browser, or recorded HTTP — all interchangeable"),
        (1, "PrimeFaces polling is filtered out automatically"),
        (0, "App-server / JVM metrics need zero new Java"),
        (1, "the EAP / WildFly MicroProfile Metrics subsystem already exposes them"),
        (0, "Collector lifetime = run lifetime  →  the app is quiet the rest of the time"),
    ])

    bullets_slide(prs, "What we're NOT doing", [
        (0, "Not production monitoring — no alerting, no on-call, no SLOs"),
        (0, "Not replacing the incumbent APM"),
        (0, "Not installing a second agent on the app servers"),
        (0, "Not hand-writing MBean-reading Java (the thing that caused the JaCoCo pain last time)"),
        (0, "Not a big-bang change — the module goes into one WAR at a time"),
    ])

    bullets_slide(prs, "If we adopt this", [
        (0, "Phase in: the module into the main WAR + 3–4 user actions first"),
        (0, "Reuse the existing Jenkins load-test jobs — add the run_id / release tags"),
        (0, "Point a Collector at the Test environment during a scheduled run"),
        (0, "Share the Grafana dashboards with the apps-mgmt team as a talking point"),
        (0, "Revisit EAP 8 / Jakarta EE 10 when the migration date is known"),
    ], note="Everything here is in the public repo; nothing workplace-specific is committed.")

    title_slide(prs, "Questions?",
                "Repo, dashboards, and run scripts are all in git — clone and try it",
                "github.com/mmcnicol/jboss-prometheus-grafana")

    out = HERE / "workshop-demo.pptx"
    prs.save(out)
    return out


# ---------------------------------------------------------------- deck 2
def build_deep_dive_deck():
    prs = new_deck()
    title_slide(prs, "Load-Test Metrics — How It Works & How to Adopt",
                "Architecture, the decisions behind it, and the path into the real application",
                "Technical deep-dive  ·  companion to the workshop demo")

    bullets_slide(prs, "Component map", [
        (0, "metrics-support/  —  the portable part (no demo-app dependency)"),
        (1, "metrics-api — Metrics / ActionTimer / EndpointTimer facade + the toggle + no-op"),
        (1, "metrics-micrometer / metrics-prometheus — the two backends (SPI-selected)"),
        (1, "metrics-servlet — the /metrics endpoint (web-fragment, shared)"),
        (1, "metrics-jaxrs — JAX-RS request-timing filter for services"),
        (0, "demo-app/  —  portal-web (JSF) + service-a + service-b (JAX-RS)"),
        (0, "observability/  —  Prometheus + Grafana + OTel Collector (compose)"),
        (0, "load/  —  scenario definitions + drivers (Selenium, k6 browser, k6 HTTP)"),
    ])

    bullets_slide(prs, "The toggle (requirement FR2/FR3)", [
        (0, "System property portal.metrics.enabled, default false"),
        (1, "read once at startup and cached — a restart flips it (fine for load-test envs)"),
        (1, "set in standalone.conf / JAVA_OPTS or <system-property> in standalone.xml"),
        (0, "When off:"),
        (1, "MetricsProvider returns the no-op — no registry, no meters registered"),
        (1, "the /metrics servlet returns 404 (Metrics.scrape() is empty)"),
        (1, "the JSF PhaseListener still runs but records into the no-op timer"),
        (0, "Overhead when off is unmeasurable; when on, < ~1 ms per instrumented request"),
    ])

    diagram_slide(prs, "Instrumentation — UI actions", [
        "JSF request  —  PhaseListener starts a timer at RESTORE_VIEW",
        "Action name resolved server-side: javax.faces.source / postback button / view id",
        "Polling (p:poll) requests dropped; login/save redirects handled",
        "Stop at end of RENDER_RESPONSE (or INVOKE_APPLICATION if already redirected)",
        "Metrics.get().action(name).record(elapsed, outcome)  →  portal_user_action_seconds",
    ], caption="No driver cooperation. Any driver that performs the steps gets correct metrics.")

    bullets_slide(prs, "Run labelling (Spike C)", [
        (0, "Labels applied OUTSIDE the app — in the Collector's resource processor"),
        (1, "run_id, release, test_type, env come from env vars the CI job sets per run"),
        (1, "the same app build serves every run"),
        (0, "Collector = one process, lifetime = one run"),
        (1, "not in the always-on stack; a compose 'loadtest' profile"),
        (1, "scrapes the app + WildFly :9990, prometheusremotewrite to Prometheus"),
        (0, "Pushgateway rejected (wrong tool); k6 remote-write kept for the client-side view"),
    ])

    bullets_slide(prs, "Dashboards", [
        (0, "User Actions — Load Test  —  per-action rate / error% / p50-p99"),
        (0, "Baseline vs Candidate  —  bar + delta table (alignment-free) + offset overlay"),
        (0, "Release Trend  —  one point per run over weeks, backed by recording rules"),
        (0, "Service Endpoints  —  per-endpoint p95, client (k6) vs server"),
        (0, "App Server & JVM  —  heap / GC / threads / CPU / Undertow, from MP Metrics"),
        (1, "all provisioned from checked-in JSON — reproducible, no click-ops"),
    ])

    bullets_slide(prs, "Spike A — Micrometer vs Prometheus client", [
        (0, "Both built behind one MetricsBackend SPI; compared on a real deploy"),
        (0, "Similar transitive footprint (Micrometer 1.13 pulls the same client_java 1.x)"),
        (0, "Micrometer wins on ad-hoc counter tags and the binder ecosystem"),
        (0, "Recommendation: Micrometer default, aligns with a future Spring Boot direction"),
        (0, "Prometheus client stays wired — proof the facade is backend-neutral"),
    ])

    bullets_slide(prs, "Spike E — app-server / JVM metrics", [
        (0, "Q: do we need a 2nd agent, or hand-write MBean code?  A: neither."),
        (0, "WildFly / EAP MicroProfile Metrics subsystem exposes JVM + server metrics"),
        (1, "already in the default EAP 7.4 profile; unauthenticated on the mgmt port"),
        (1, "Collector scrapes :9990/metrics during the run — nothing installed"),
        (0, "jmx_exporter as a -javaagent conflicts with WildFly boot (LogManager ordering)"),
        (1, "standalone jmx_exporter over remote JMX works, if the subsystem isn't available"),
        (0, "Nothing here needs a unit test or a coverage exclusion"),
    ])

    bullets_slide(prs, "Decisions on record", [
        (0, "D1 free PrimeFaces Saga theme + a documented swap point"),
        (0, "D3 release label = git describe, normalised to a bare version"),
        (0, "D4 recording rules only — no separate results store"),
        (0, "D5 JDK 17 (matches EAP 7.4 at the workplace)"),
        (0, "D6 Selenium/Java is the primary UI driver; k6 browser the alternative"),
        (0, "Full list: docs/01-requirements.md §8; rationale in docs/findings/"),
    ])

    diagram_slide(prs, "Adoption path", [
        "1.  Add metrics-api + metrics-micrometer + metrics-servlet to the main WAR",
        "2.  Register the PhaseListener; instrument 3–4 top user actions",
        "3.  Add run_id / release tags to the existing Jenkins load-test jobs",
        "4.  Run a Collector against Test during a scheduled load test",
        "5.  Roll the JAX-RS filter into services that already have API tests",
        "6.  Enable the MicroProfile Metrics scrape for JVM / app-server view",
    ], caption="Each step is independently shippable and reversible.")

    bullets_slide(prs, "Risks & mitigations", [
        (0, "MP Metrics subsystem stripped in a hardened EAP profile → standalone jmx_exporter"),
        (0, "Management port exposure → keep internal / firewalled (NFR7)"),
        (0, "Per-run run_id cardinality → few runs; recording rules collapse each run"),
        (0, "JSF action mis-attribution → covered by tests; Spike B write-up"),
        (0, "EAP 8 / Jakarta EE 10 → javax→jakarta on the module; desk review pending"),
    ])

    title_slide(prs, "Appendix", "docs/ has the requirements, the spikes, the technical solution, and every finding",
                "github.com/mmcnicol/jboss-prometheus-grafana")

    out = HERE / "technical-deep-dive.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    for f in (build_demo_deck(), build_deep_dive_deck()):
        print("wrote", f)
